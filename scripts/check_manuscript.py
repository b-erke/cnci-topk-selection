#!/usr/bin/env python3
"""Check a manuscript .docx against the replication outputs: every cell of Tables 1–6 and a list of reported numbers.
    python scripts/check_manuscript.py --docx manuscript.docx --results results
Exit code 1 if any check fails. Requires python-docx."""
import argparse, sys
from pathlib import Path
import pandas as pd

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--docx", type=Path, required=True); ap.add_argument("--results", type=Path, default=Path("results"))
    a = ap.parse_args(); R = a.results
    import docx
    doc = docx.Document(a.docx)
    text = "\n".join(p.text for p in doc.paragraphs)
    T = [[[c.text.strip() for c in r.cells] for r in t.rows] for t in doc.tables]
    rd = lambda f: pd.read_csv(R / f)
    fx, ref, fits, bnd, prof = rd("table1_fixed_k.csv"), rd("copula_references.csv"), rd("copula_fits.csv"), rd("table4_boundary.csv"), rd("table5_profiles.csv")
    ho, hv, ag, bt, thr, reg = rd("hierarchy_overlap.csv").query("map == 'main'"), rd("hierarchy_variants.csv").query("map == 'main'"), rd("agreement.csv").set_index("level"), rd("bootstrap_coef_diff.csv").set_index("level"), rd("threshold_sensitivity.csv"), rd("regressions.csv")
    fails, n = [], 0
    def check(ok, what):
        nonlocal n; n += 1
        if not ok: fails.append(what)
    C, O = "Country/region", "Organization"
    # Table 1
    for r in T[0][1:]:
        x = fx[(fx.level == r[0]) & (fx.K == int(r[1]))].iloc[0]
        check(r[2:] == [str(int(x.shared)), f"{x.turnover*100:.1f}%", f"{x.jaccard:.3f}"], f"Table 1 {r}")
    # Table 2
    for r in T[1][1:]:
        x = ref[(ref.family == "gauss_cal") & (ref.level == r[0]) & (ref.K == int(r[1]))].iloc[0]
        check(r[2:] == [str(int(x.observed)), f"{x['mean']:.2f}", f"{int(x.p05)}–{int(x.p95)}", f"{x.p_sim_ge_obs:.3f}"], f"Table 2 {r}")
    # Table 3
    fam = {"Gaussian": "gauss", "t, ν = 10": "t10", "t, ν = 5": "t5", "t, ν = 3": "t3", "Gumbel": "gumbel", "Survival Clayton": "sclayton",
           "Frank": "frank", "Clayton": "clayton", "Survival BB1": "sbb1", "t, ν estimated": "tfree"}
    for r in T[2][2:]:
        f = fam[r[0]]; fo = fits[(fits.level == O) & (fits.family == f)].iloc[0]; fc = fits[(fits.level == C) & (fits.family == f)].iloc[0]
        check(r[3:5] == [f"{fc.delta_pseudo_aic:.1f}", f"{fo.delta_pseudo_aic:.1f}"], f"Table 3 pseudo-AIC {r[0]}")
        if f != "tfree":
            check(r[1:3] == [f"{fo.lambda_L_cal:.2f}", f"{fo.lambda_U_cal:.2f}"], f"Table 3 tail dependence {r[0]}")
            for cell, K in zip(r[5:], (20, 50, 100)):
                x = ref[(ref.family == f + "_cal") & (ref.level == O) & (ref.K == K)].iloc[0]
                check(cell == f"{x['mean']:.1f} ({int(x.p05)}–{int(x.p95)})", f"Table 3 {r[0]} K={K}")
    # Table 4
    for r in T[3][1:]:
        x = bnd[(bnd.level == r[0]) & (bnd.K == int(r[1]))].iloc[0]
        check(r[2:] == [str(int(x.changed_units)), f"{x.median_rank_gap:.1f}", f"{x.median_rank_gap_over_N:.3f}", f"{x.median_iqr_gap:.3f}"], f"Table 4 {r}")
    # Table 5
    for r in T[4][1:]:
        x = prof[(prof.level == r[0]) & (prof.K == int(r[1])) & (prof.group == r[2])].iloc[0]
        check(r[3:] == [str(int(x.n)), f"{x.mean_intl:.2f}", f"{x.mean_quad:.2f}"], f"Table 5 {r}")
    # Table 6
    for var, r in zip(hv.variant, T[5][1:]):
        check(r[1:3] == [f"{int(hv[hv.variant == var].N.iloc[0]):,}", f"{hv[hv.variant == var].spearman.iloc[0]:.3f}"], f"Table 6 N/rho {var}")
        for cell, K in zip(r[3:], (10, 20, 50, 100, 200)):
            g = ho[(ho.variant == var) & (ho.K == K) & (ho.family == "gauss_cal")].iloc[0]; b = ho[(ho.variant == var) & (ho.K == K) & (ho.family == "sbb1_cal")].iloc[0]
            check(cell == f"{int(g.observed)} ({g.p_sim_ge_obs:.3f}; {b.p_sim_ge_obs:.3f})", f"Table 6 {var} K={K}")
    # reported numbers in the text
    t = thr.set_index(["level", "min_docs"]).spearman
    q = lambda lev, model, sample, term, col: reg[(reg.level == lev) & (reg.model == model) & (reg["sample"] == sample) & (reg.term == term)][col].iloc[0]
    expected = [f"{ag.loc[C,'spearman']:.4f}", f"{ag.loc[O,'spearman']:.4f}", f"{ag.loc[C,'kendall']:.4f}", f"{ag.loc[O,'kendall']:.4f}",
                f"{ag.loc[C,'rbo_ext_p0.90']:.3f}", f"{ag.loc[O,'rbo_ext_p0.90']:.3f}", f"{ag.loc[C,'rbo_ext_p0.98']:.3f}", f"{ag.loc[O,'rbo_ext_p0.98']:.3f}",
                f"{bt.loc[C,'point']:.4f}", f"{bt.loc[C,'ci_low']:.4f}–{bt.loc[C,'ci_high']:.4f}", f"{bt.loc[O,'point']:.4f}", f"{bt.loc[O,'ci_low']:.4f}–{bt.loc[O,'ci_high']:.4f}",
                *[f"{t[(lev, k)]:.3f}" for lev in (C, O) for k in (50, 100, 200, 500)],
                f"{q(C,'diff ~ shares + log docs','primary','quad','r2'):.3f}", f"{q(O,'diff ~ shares + log docs','primary','quad','r2'):.3f}"]
    gp = R / "copula_gof.csv"
    if gp.exists():   # goodness-of-fit values reported in Sections 2.4 and 3.3
        gf = pd.read_csv(gp); g = lambda lev, f, col, qq=0.8: gf[(gf.level == lev) & (gf.family == f) & (gf.q_upper == qq)][col].iloc[0]
        tp = [g(O, f, "p_S_n_upper", 0.9) for f in ("t10", "t5", "t3", "tfree")]
        expected += [f"{int(gf.B.iloc[0]):,} replicates", f"p = {g(O,'sbb1','p_S_n'):.3f})", f"p = {g(O,'sbb1','p_S_n_upper',0.8):.3f} for q = 0.8",
                     f"p = {g(O,'sbb1','p_S_n_upper',0.9):.3f} for q = 0.9", f"p = {g(O,'gauss','p_S_n_upper',0.9):.3f})",
                     f"p = {min(tp):.3f}–{max(tp):.3f}", f"{int(g(O,'sbb1','n_upper',0.8))} organizations", f"{int(g(O,'sbb1','n_upper',0.9))} organizations",
                     f"p ≥ {min(g(C,f,'p_S_n') for f in ('gauss','t10','t5','t3','gumbel','sbb1','tfree')):.2f}",
                     f"p ≤ {max(g(C,f,'p_S_n') for f in ('sclayton','frank','clayton')):.3f}"]
    for e in expected:
        check(e in text, f"text value {e} not found")
    print(f"{n - len(fails)} of {n} checks passed")
    for f in fails: print("FAIL:", f)
    sys.exit(1 if fails else 0)

if __name__ == "__main__":
    main()
